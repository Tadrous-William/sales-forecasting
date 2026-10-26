"""Step 3: Excel forecast workbook (stock plan, weekly forecast, accuracy) and portfolio charts."""
import json
from pathlib import Path

import matplotlib
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

matplotlib.use("Agg")
import matplotlib.dates as mdates  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
DATA, OUT, SHOW = ROOT / "data", ROOT / "output", ROOT / "portfolio"
SHOW.mkdir(exist_ok=True)
INK, INK2, MUTED, GRID = "#0b0b0b", "#52514e", "#8a8983", "#e7e6e1"
BLUE, BLUE_LIGHT, GREY_BAR = "#2a78d6", "#cfe0f6", "#c9c8c2"
NAMES = {"last_week": "Last week", "avg_8w": "8-week average", "avg_8w_per_day": "8-week average per trading day",
         "same_week_ly": "Same week last year", "ly_x_trend": "Last year x trend", "lightgbm": "ML model (median)",
         "lightgbm_mean": "ML model (average)", "combined": "Combined forecast (used)"}
plt.rcParams.update({"font.family": "Segoe UI", "axes.edgecolor": GRID, "axes.labelcolor": INK2, "xtick.color": INK2,
                     "ytick.color": INK2, "axes.spines.top": False, "axes.spines.right": False})


def stock_plan(fc: pd.DataFrame, products: pd.DataFrame, wp: pd.DataFrame) -> pd.DataFrame:
    """Next 4 weeks: expected demand and the stock that covers demand 9 times out of 10."""
    first4 = fc[fc["week"] <= fc["week"].min() + pd.Timedelta(weeks=3)]
    g = first4.groupby("stock_code")
    plan = pd.DataFrame({"demand_4w": g["forecast"].sum(),
                         "stock_90": g["forecast"].sum() + np.sqrt(g.apply(lambda d: ((d["high"] - d["forecast"]) ** 2).sum()))})
    plan["demand_8w"] = fc.groupby("stock_code")["forecast"].sum()
    recent = wp[wp["week"] > wp["week"].max() - pd.Timedelta(weeks=8)].groupby("stock_code")["units"].mean()
    plan["avg_week_last_8"] = recent
    plan = products.set_index("stock_code")[["product_name", "category", "revenue_last_52w"]].join(plan).reset_index()
    return plan.sort_values("revenue_last_52w", ascending=False)


def write_excel(plan, fc, products, metrics, path):
    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0F172A")

    def sheet(ws, cols, rows, fmts=None):
        for c, (h, w) in enumerate(cols, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font, cell.fill, cell.alignment = head_font, head_fill, Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.row_dimensions[1].height = 50
        for r, row in enumerate(rows, 2):
            for c, v in enumerate(row, 1):
                cell = ws.cell(row=r, column=c, value=v)
                if fmts and fmts.get(c):
                    cell.number_format = fmts[c]
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"

    ws = wb.active
    ws.title = "Stock plan"
    weeks = sorted(fc["week"].unique())
    sheet(ws, [("Stock code", 11), ("Product", 40), ("Category", 20), ("Sold per week, last 8 weeks", 14),
               ("Expected demand, next 4 weeks", 14), ("Stock to cover 4 weeks (90% of the time)", 16),
               ("Expected demand, next 8 weeks", 14), ("Revenue last 52 weeks (£)", 14)],
          [[r.stock_code, r.product_name, r.category, round(r.avg_week_last_8), round(r.demand_4w), round(r.stock_90),
            round(r.demand_8w), round(r.revenue_last_52w)] for r in plan.itertuples()],
          {4: "#,##0", 5: "#,##0", 6: "#,##0", 7: "#,##0", 8: "#,##0"})
    for r in range(2, len(plan) + 2):
        ws.cell(row=r, column=6).fill = PatternFill("solid", fgColor="DBEAFE")

    ws = wb.create_sheet("Weekly forecast")
    names = products.set_index("stock_code")["product_name"]
    wide = fc.pivot(index="stock_code", columns="week", values="forecast").reindex(plan["stock_code"])
    sheet(ws, [("Stock code", 11), ("Product", 40)] + [(f"Week of {pd.Timestamp(w):%d %b %Y}", 12) for w in weeks],
          [[code, names[code]] + [round(v) for v in wide.loc[code]] for code in wide.index],
          {c: "#,##0" for c in range(3, 3 + len(weeks))})
    ws = wb.create_sheet("Low - high range")
    lo = fc.pivot(index="stock_code", columns="week", values="low").reindex(plan["stock_code"])
    hi = fc.pivot(index="stock_code", columns="week", values="high").reindex(plan["stock_code"])
    sheet(ws, [("Stock code", 11), ("Product", 40)] + [(f"Week of {pd.Timestamp(w):%d %b}", 14) for w in weeks],
          [[code, names[code]] + [f"{lo.loc[code, w]:,.0f} - {hi.loc[code, w]:,.0f}" for w in weeks] for code in lo.index])

    ws = wb.create_sheet("Accuracy test")
    lv = [("product_week", "One product, one week"), ("product_4_weeks", "One product, next 4 weeks"),
          ("store_week", "All 200 products, one week"), ("store_4_weeks", "All 200 products, next 4 weeks")]
    sheet(ws, [("Method", 34)] + [(label + " (error)", 16) for _, label in lv] + [("Bias", 10)],
          [[NAMES[m]] + [metrics["wape"][k][m] for k, _ in lv] + [metrics["bias"][m]] for m in NAMES],
          {c: "0.0%" for c in range(2, 7)})
    note = ws.cell(row=len(NAMES) + 3, column=1, value=(
        "Error = total absolute difference between forecast and actual units / actual units (WAPE). Tested on 7 forecast dates "
        f"in 2011 ({metrics['product_weeks_tested']:,} product-weeks); each forecast only used data before its date."))
    note.alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=len(NAMES) + 3, start_column=1, end_row=len(NAMES) + 5, end_column=6)
    wb.save(path)


def chart_store_backtest(bt: pd.DataFrame, total_units: pd.Series, path: Path):
    st = bt.groupby(["origin", "week"])[["combined", "actual", "low", "high"]].sum().reset_index()
    fig, ax = plt.subplots(figsize=(12, 4.8), dpi=160)
    hist = total_units[(total_units.index >= "2011-01-01")]
    top200 = bt.groupby("week")["actual"].first()  # not used; actuals come from the forecast rows below
    act = st.groupby("week")["actual"].first()
    ax.plot(act.index, act.values, color=INK, lw=2, label="Actual units sold (200 products)")
    for i, (o, d) in enumerate(st.groupby("origin")):
        ax.plot(d["week"], d["combined"], color=BLUE, lw=2, label="Forecast made 8 weeks ahead" if i == 0 else None)
        ax.axvline(o, color=GRID, lw=1, zorder=0)
    ax.set_ylim(0, None)
    ax.grid(axis="y", color=GRID, lw=0.8)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f"{v / 1000:.0f}k"))
    ax.set_title("Weekly units: forecasts made on 7 dates in 2011 vs what really sold", loc="left",
                 fontsize=13, fontweight="bold", color=INK, pad=34)
    ax.legend(frameon=False, loc="lower left", bbox_to_anchor=(0, 1.0), ncol=2, fontsize=10)
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%b %Y"))
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def chart_accuracy(metrics: dict, path: Path):
    order = ["last_week", "same_week_ly", "ly_x_trend", "avg_8w", "avg_8w_per_day", "combined"]
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.6), dpi=160)
    for ax, (key, title) in zip(axes, [("product_4_weeks", "Each product, next 4 weeks"), ("store_4_weeks", "Whole store, next 4 weeks")]):
        vals = [metrics["wape"][key][m] for m in order]
        colors = [BLUE if m == "combined" else GREY_BAR for m in order]
        bars = ax.barh([NAMES[m] for m in order], vals, color=colors, height=0.6)
        for b, v in zip(bars, vals):
            ax.text(b.get_width() + 0.005, b.get_y() + b.get_height() / 2, f"{v:.1%}", va="center", fontsize=10, color=INK)
        ax.set_xlim(0, max(vals) * 1.2)
        ax.set_xticks([])
        ax.spines["bottom"].set_visible(False)
        ax.tick_params(axis="y", length=0, labelsize=10)
        ax.set_title(title, loc="left", fontsize=12, fontweight="bold", color=INK)
    fig.suptitle("Forecast error (lower is better), tested on 7 dates in 2011", x=0.01, ha="left", fontsize=14,
                 fontweight="bold", color=INK)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def chart_product(code, name, wp, fc, bt, path):
    h = wp[wp["stock_code"] == code].set_index("week")["units"]
    h = h[h.index >= "2011-01-01"]
    f = fc[fc["stock_code"] == code].set_index("week")
    fig, ax = plt.subplots(figsize=(12, 4.4), dpi=160)
    ax.plot(h.index, h.values, color=INK, lw=1.6, label="Units sold per week")
    ax.fill_between(f.index, f["low"], f["high"], color=BLUE_LIGHT, lw=0, label="Likely range (8 in 10 weeks)")
    ax.plot(f.index, f["forecast"], color=BLUE, lw=2.2, marker="o", ms=4, label="Forecast, next 8 weeks")
    ax.set_ylim(0, None)
    ax.grid(axis="y", color=GRID, lw=0.8)
    ax.set_title(f"{name.title()} ({code}): weekly units and the next 8 weeks", loc="left", fontsize=13,
                 fontweight="bold", color=INK, pad=34)
    ax.legend(frameon=False, loc="lower left", bbox_to_anchor=(0, 1.0), ncol=3, fontsize=10)
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%b %Y"))
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


if __name__ == "__main__":
    metrics = json.loads((OUT / "metrics.json").read_text())
    products = pd.read_parquet(DATA / "products.parquet")
    wp = pd.read_parquet(DATA / "weekly_product.parquet")
    total = pd.read_parquet(DATA / "weekly_total.parquet").set_index("week")["units"]
    fc = pd.read_parquet(OUT / "forecast.parquet")
    bt = pd.read_parquet(OUT / "backtest.parquet")
    plan = stock_plan(fc, products, wp)
    plan.to_csv(OUT / "stock_plan.csv", index=False)
    write_excel(plan, fc, products, metrics, OUT / "forecast_next_8_weeks.xlsx")
    chart_store_backtest(bt, total, SHOW / "backtest_store.png")
    chart_accuracy(metrics, SHOW / "accuracy.png")
    for code in plan["stock_code"].head(3):
        chart_product(code, products.set_index("stock_code").loc[code, "product_name"], wp, fc, bt, SHOW / f"product_{code}.png")
    print(plan.head(8)[["product_name", "avg_week_last_8", "demand_4w", "stock_90"]].round(0).to_string(index=False))
